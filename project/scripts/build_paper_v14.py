"""Create the three-page paper from retained evidence and v0.14 analytical outputs."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];O=ROOT/'deliverables'
s=json.loads((ROOT/'data/visualization_summary_v13.json').read_text())
d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor.from_string('192C42')
d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(6)
d.styles['Normal'].paragraph_format.line_spacing=1.05
d.styles['Title'].font.size=Pt(23)
for style in d.styles:
    for border in style.element.xpath('.//w:pBdr'):border.getparent().remove(border)
for name in ['Heading 1','Heading 2']:
    d.styles[name].font.size=Pt(13);d.styles[name].paragraph_format.space_before=Pt(8);d.styles[name].paragraph_format.space_after=Pt(5)
f=sec.footer.paragraphs[0];f.add_run('15 October 2026  |  Informative research  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
for r in f.runs:r.font.size=Pt(8)
def p(t):d.add_paragraph(t)
def h(t):d.add_heading(t,1)
def table(headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers));t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,x in zip(t.rows[0].cells,headers):c.text=x
    for row in rows:
        for c,x in zip(t.add_row().cells,row):c.text=x
    for i,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'192C42' if i==0 else 'F1F5F8');c._tc.get_or_add_tcPr().append(shade)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(4);pp.paragraph_format.space_before=Pt(4)
                for r in pp.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '192C42')
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))

d.add_heading('The state of U.S. trade facilitation programs',0)
d.add_paragraph('Instruments, geographic reach and evidence of operation',style='Subtitle')
p('U.S. trade facilitation combines commercial matchmaking, export finance, project preparation, investment finance and assistance with the legal and institutional conditions for trade. The evidence supports an informative account of these functions and selected country activity. It does not establish a complete census of programs currently delivering services or an official roster of “Trade over Aid” participant countries. Unresolved legacy awards can be isolated without preventing analysis of the broader system. [1]')
p('The research register contains 124 screened records within a defined 22-agency frame: 62 direct instruments, 20 supporting services, 17 overlapping initiatives, 11 unresolved candidates, three adjacent routes, ten exclusions and one alias. These are heterogeneous counting units, not 124 independent statutory programs. Named initiatives and their component awards are retained separately to avoid treating multiple administrative layers as additional deals. [1]')
h('How the instruments facilitate trade')
table(['Function','Illustrative instruments','Scope and geographic constraint'],[
('Finding and preparing transactions','Commerce partner search, Gold Key, advocacy; USTDA studies and pilots','Local service capacity, designated events or selected projects; an overseas office does not guarantee every service.'),
('Financing exports and investment','EXIM insurance, loans and guarantees; DFC debt, guarantees, equity and insurance; SBA export finance','Product-specific credit, borrower, project and destination conditions. Financing commitments differ from disbursements.'),
('Developing markets and enabling infrastructure','USDA/FAS promotion and credit instruments; CLDP legal assistance; selected State and MCC activities','Agriculture, energy, digital systems, infrastructure or institutional reform, depending on the instrument and award.')],[1.45,2.6,3.05])
p('Sector specialization matters. USTDA preparation and pilot instruments address critical minerals, digital infrastructure, energy and transportation; EXIM’s Supply Chain Resiliency Initiative focuses on critical minerals and rare earths. USDA’s GSM-102 concerns agricultural exports, State’s FIRST addresses civil nuclear small modular reactors, and Foreign Military Sales covers defense articles, services and training. Broad Commerce services sit alongside these specialized routes. [1]')
p('Commerce, EXIM and DFC retain official product descriptions. CLDP identifies itself within Commerce’s Office of General Counsel and describes interagency funding. These sources establish functions, but not each legacy agreement, intake capacity or country delivery. [2–5]')
h('Evidence of operation is uneven')
p('Across the register, nine records have recent delivery reports, 12 have portfolio or authorization evidence, and 14 have funding, intake or approval actions as their strongest evidence. Another 60 have an official offer only, 18 remain operationally unresolved, and 11 are exclusions or aliases. The first three categories must not be added and described as independently verified active programs: an approval, a financing stock and delivered assistance are different observations. [1]')

d.add_page_break();h('Geographic reach: activity, eligibility and offices')
p('The country ledger contains 2,183 source observations. Its qualifying activity matrix contains 281 country–program pairs across 19 inventory records and 124 country or territory units. The full frame covers 249 units and 110 records. Published financing datasets dominate this coverage, so the map shows where this collection contains evidence, not the intensity or complete distribution of U.S. trade facilitation. [1, 6]')
p('The service review adds CLDP delivery in Algeria, Montenegro, Sweden and Tunisia, plus Advocacy Center support for an Air Cambodia transaction reported February 3, 2026. The latter identifies the buyer’s country, not the signing venue or aircraft delivery. Historical Gold Key cases and recipient-country-only training records remain separate. These observations do not establish service effectiveness or comprehensive geographic coverage. [1]')
table(['Country evidence','Meaning in the atlas','Interpretation limit'],[
('Reported delivery (D)','A dated primary account reports assistance or an activity','Not independently measured impact or proof of continuing capacity.'),
('Financing / implementation (A, C, I)','EXIM authorization; DFC commitment or signed financing; MCC implementation roster','Not interchangeable with disbursement, construction completion or a completed trade deal.'),
('Pipeline / targets; historical','Approvals, announcements, award targets, closed or terminated projects','Preserved in the full status data; excluded from the qualifying activity count.'),
('Eligibility only; unresolved','A conditional geographic route, or insufficient evidence','Eligibility is not delivery. Unknown is not inactive or ineligible.')],[1.5,2.8,2.8])
p('Geography itself has several meanings: EXIM credit-risk location, DFC project location, an assistance venue, an award’s target market and a participant’s home country. The ledger preserves these roles. Participant-origin records are not automatically converted into overseas delivery. Financial approvals and commitments retain their source dates, including the June 30, 2026 portfolio cutoff. [6]')
p('The retained directory evidence identifies published in-country network addresses in 67 country units for the Commercial Service and 78 for FAS. Explicit embassy or mission affiliation is recorded in 64 and 68 units respectively; these categories overlap and are not additive. These are deduplicated agency-network observations, not current staff counts or assignments of each program to each embassy. Other agencies retain weaker or unresolved presence evidence. [1]')
h('Comparing geographic rules and scope')
p('A categorical two-dimensional chart displays all 62 direct instruments by geographic rule and broad scope. A separate exploratory similarity plot includes 13 instruments with at least 70% resolved foreign-country route entries, yielding seven distinct coded profiles. It combines geographic-route mismatch and scope mismatch at equal weights; unknown country entries are omitted pairwise. The 49 excluded instruments remain in the inventory and categorical chart. [1]')
p('The numerical plot is a comparison of eligibility architecture, not actual country delivery. Its normalized two-dimensional stress is 0.215, indicating material projection distortion; exact distances are therefore supplied alongside it. Changing geography’s weight from 25% to 75% preserves the ordering of pairwise distances in this small subset, but does not validate the coding or make the instruments interchangeable. Similar DFC profiles reflect shared coarse geography and scope, not identical products. [1]')

d.add_page_break();h('Legacy awards: uncertainty stays local')
p('The targeted legacy review covers 25 awards linked or provisionally linked to Prosper Africa, Power Africa and digital-connectivity initiatives. None has certified current survival in this review; this is not evidence that none survive. Historical notices, later administrative actions, current ownership and country delivery remain distinct questions. The atlas also displays the two unresolved CLDP/DFC interagency-agreement leads. [1, 7]')
p('The uploaded USAspending exports corroborate ten prime-award transactions across Green Powered Technology (72067424F00006), RTI (72067423C00003) and Chemonics (72067424C00001). Chemonics’ March 7, 2025 P00003 is a zero-dollar administrative modification with a truncated description; it does not establish reinstatement. RTI’s export additionally reports a $40,250 subaward action to Bowman Gilfillan dated April 15, 2025, after the prime’s reported termination notice. Its country field is blank. This is a dated subaward-reporting lead, not sufficient evidence of prime reinstatement or current country delivery. [7–8]')
p('The CLDP agreement 72067424T00001 and DFC PAPA 72067420T00001 remain exact-ID leads from a mirrored legacy roster. Their executed agreements and subsequent amendments are not authenticated. This does not invalidate separately documented CLDP activity or DFC financing; evidence for those broader programs must not be assigned to these specific instruments without a verified link. [4–5, 8–9]')
h('Use and limits')
p('The paper and atlas can inform understanding of the U.S. instrument landscape, documented geographic reach, sector scope, overseas networks and evidence gaps. They cannot establish program effectiveness, additional exports caused by assistance, current staffing, universal intake availability or a definitive active-program total. The remaining evidence gaps are attached to the affected records rather than used as a reason to withhold the whole descriptive product.')
p('The next research priorities are to reconcile selected post-notice action files, authenticate the two interagency agreements, and expand comparable country-delivery records for service programs. These should improve later editions without delaying use of the present qualified account. The full package retains source snapshots, dates, identifiers, exclusions, reproducible scripts and the distinctions between observations, eligibility and inferred similarity.')
h('Selected sources and reproducibility')
sources=[
'[1] Companion data: program_census.csv; country_activity_ledger.csv; country_program_status_v14.csv; network_evidence_summary_v14.csv; similarity distances, exclusions and sensitivity; source_registry_v11.csv. Methods and full URLs accompany each record.',
'[2] Commerce, Gold Key Service: https://www.trade.gov/gold-key-service',
'[3] EXIM, Solutions: https://www.exim.gov/solutions',
'[4] DFC, Asset Classes: https://www.dfc.gov/what-we-offer/asset-classes',
'[5] Commerce CLDP, FAQ: https://cldp.doc.gov/about-cldp/faq',
'[6] EXIM FY2026 Q3 data: https://img.exim.gov/s3fs-public/dataset/vbhv-d8am/data-gov_fy26-q3.csv ; DFC: https://www3.dfc.gov/OPICProjects/Data/ActiveProjects.xlsx ; MCC: https://www.mcc.gov/where-we-work/',
'[7] DOJ, Global Health Council v. Trump, ECF 20-1, filed February 12, 2025, pp. 2 and 11. RECAP document and government cover filing retained in research_v11/evidence.',
'[8] User-supplied USAspending contract exports dated October 15, 2026, for the three IDs above. TransactionHistory and Sub-Awards CSVs retained unchanged under research_v13/uploaded_records.',
'[9] Mirrored USAID March 6, 2025 roster, PDF pp. 277 and 283: https://i-base.info/wp-content/uploads/2025/03/USAID-Terminated-Awards_2025-03-06_Clean.pdf . Qualified provenance; not a signed agreement.'
]
for t in sources:
    pp=d.add_paragraph(t);pp.paragraph_format.space_after=Pt(3)
    for r in pp.runs:r.font.size=Pt(8)
d.save(O/'US_Trade_Facilitation_Informative_Paper_v0_14.docx')
print('Paper written')
