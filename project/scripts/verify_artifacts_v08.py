from pathlib import Path
import csv,json,datetime,fitz
from openpyxl import load_workbook
P=Path(__file__).resolve().parents[1]
mapping={'Claim audit':'claim_audit','Corrections':'corrections','Legacy review':'legacy_reconciliation','Selected awards':'award_evidence','Award actions':'award_actions','Budget evidence':'budget_evidence','Inventory':'inventory_view','Ownership':'ownership_review','Changes':'inventory_changes','Eligibility':'eligibility_matrix','Scope profiles':'program_scope_profiles','Geographic rules':'program_geographic_rules','MCC forecast':'mcc_procurement_forecast','Evidence gaps':'remaining_evidence_gaps','Sources':'expansion_sources','Codebook':'codebook','Country activity':'country_program_activity_matrix','Country stages':'country_program_stage_matrix','DFC unallocated':'dfc_unallocated_activity_matrix','Activity ledger':'country_activity_ledger','Activity coverage':'activity_coverage','EXIM FY2026':'exim_fy2026_authorizations','DFC portfolio':'dfc_active_portfolio','MCC lifecycle':'mcc_lifecycle','MDCP targets':'mdcp_country_targets','Current sources':'source_registry_v08','Activity codes':'activity_codebook','Agency frame':'agency_coverage_screen'}
w=load_workbook(P/'deliverables/US_Trade_Facilitation_Inventory_v0_8.xlsx',read_only=True,data_only=True);checked=0
assert len(w.sheetnames)==29
for sheet,key in mapping.items():
 expected=list(csv.reader(open(P/'data'/(key+'.csv'),encoding='utf-8-sig')))
 actual_rows=list(w[sheet].values)
 assert len(actual_rows)==len(expected),(sheet,len(actual_rows),len(expected))
 for i,(actual,row) in enumerate(zip(actual_rows,expected)):
  for j,b in enumerate(row):
   a=actual[j]
   if isinstance(a,datetime.datetime):
    if 'T' in b:
     target=datetime.datetime.fromisoformat(b).replace(tzinfo=None)
     assert abs((a-target).total_seconds())<.002,(sheet,i,j,a,b)
    else:assert a.strftime('%Y-%m-%d')==b,(sheet,i,j,a,b)
   elif isinstance(a,(int,float)) and not isinstance(a,bool):assert abs(float(a)-float(b))<1e-5,(sheet,i,j,a,b)
   else:assert ('' if a is None else str(a))==b,(sheet,i,j,a,b)
   checked+=1
assert w['Read first']['B7'].value==12436
assert w['Read first']['B17'].value==275
f=fitz.open(P/'deliverables/US_Trade_Facilitation_Informative_Paper_v0_8.pdf');assert len(f)==3
txt='\n'.join(p.get_text() for p in f)
for x in ['124 screened','2,169','275 observed','121 country','908 non-cancelled','1,085','294 Finance','seven of 12']:assert x in txt,x
for x in ['itinerary','travel tips','before travel']:assert x not in txt
r={'workbook_sheets':29,'csv_cells_verified':checked,'paper_pages':3,'paper_key_counts_verified':True,'visual_review':'Paper and changed/added workbook views inspected; retained tabs reuse prior visual QA.'}
(P/'validation/artifact_verification_v08.json').write_text(json.dumps(r,indent=2));print(json.dumps(r,indent=2))
