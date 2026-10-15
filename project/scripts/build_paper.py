from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(__file__).resolve().parents[1]
O=ROOT/'deliverables'
d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.6);sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
n=d.styles['Normal'];n.font.size=Pt(10.5);n.paragraph_format.space_after=Pt(6);n.paragraph_format.line_spacing=1.05
for name,size in [('Title',23),('Heading 1',14),('Heading 2',11)]:
 d.styles[name].font.size=Pt(size);d.styles[name].paragraph_format.space_before=Pt(9);d.styles[name].paragraph_format.space_after=Pt(5)
d.styles['Caption'].font.size=Pt(9)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
f.add_run('15 October 2026  |  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
for rr in f.runs:rr.font.size=Pt(8);rr.font.color.rgb=RGBColor(85,85,85)
def para(t,boldlead=None):
 p=d.add_paragraph()
 if boldlead and t.startswith(boldlead):p.add_run(boldlead).bold=True;p.add_run(t[len(boldlead):])
 else:p.add_run(t)
 return p
def head(t):d.add_heading(t,1)
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for c,x in zip(t.rows[0].cells,headers):c.text=x
 for row in rows:
  cells=t.add_row().cells
  for c,x in zip(cells,row):c.text=x
 for i,row in enumerate(t.rows):
  for j,c in enumerate(row.cells):
   c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    z=OxmlElement('w:'+edge);z.set(qn('w:val'),'single');z.set(qn('w:sz'),'4');z.set(qn('w:color'),'D9D9D9');b.append(z)
   tcpr.append(b);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'244D69' if i==0 else ('F1F5F7' if i%2==0 else 'FFFFFF'));tcpr.append(sh)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    z=OxmlElement('w:'+edge);z.set(qn('w:w'),'75');z.set(qn('w:type'),'dxa');margins.append(z)
   tcpr.append(margins)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
    for r in p.runs:r.font.size=Pt(9.4);r.bold=(i==0);r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
 return t

import json
S=json.loads((ROOT/'data/census_summary.json').read_text())
d.add_heading('The state of US trade facilitation programs',0)
para('Instruments and verified country evidence as of 15 October 2026').style='Subtitle'
para('U.S. trade facilitation combines commercial services, export credit, project preparation, investment finance and sector assistance. The public record supports a broad account of these instruments and selected country activity. It does not support a single definitive count of all programs currently delivering services. New coordination initiatives coexist with continuing financing portfolios, closed application rounds and legacy projects affected by the foreign-assistance review.')
para('The expanded research register contains 124 screened records: 62 direct services or instruments, 20 supporting services, 17 overlapping initiatives, 11 unresolved candidates, three adjacent routes, ten exclusions and one alias. Each record now has an operational-evidence disposition. The review covers a defined 22-agency frame and targeted omitted instruments; it is not an agency-certified census. Service products and statutory programs are different counting units. [1]')
head('What the instruments do')
table(['Function','Principal instruments and sector interests'],[
('Find buyers and reduce transaction costs','Commerce partner searches, market research, trade missions, advocacy and compliance assistance serve multiple industries. USDA market development and missions concentrate on agriculture.'),
('Finance exports and investment','EXIM insurance, guarantees and loans address export credit risks. SBA export lending supports eligible U.S. businesses. DFC debt, guarantees, insurance, equity and funds support overseas investments, with current emphasis on critical minerals, infrastructure and technology.'),
('Prepare projects and institutions','USTDA funds project preparation and pilots and convenes commercial exchanges. FIRST supports civil nuclear readiness. CLDP and State sector assistance address legal and regulatory barriers. MCC compacts support country infrastructure and institutional investment.'),
('Facilitate defense trade','Foreign Military Sales and Foreign Military Financing operate through separate security-assistance authorities. Their inclusion makes this register broader than civilian export promotion.')
],[1.6,5.5])
head('Evidence of current operation remains uneven')
para('Across the 124 records, eight have recent delivery reports, 12 have current portfolio or authorization evidence, and 14 have funding, intake or approval actions as their strongest evidence. Another 61 have only a current official offer, 18 remain operationally unresolved, and 11 are exclusions or aliases. These categories describe evidence, not 34 independently verified active programs. Funding action can mean a closed competition or board approval, while a live service page cannot establish capacity.')
para('Eight additions are recorded as supporting services or overlays, leaving the 62 direct-tool count unchanged. They include Commerce’s American AI Exports Program and SABIT, three named DFC arrangements, Pax Silica, State’s Power Sector Program and the Japan–U.S.–Mekong Power Partnership. Commerce reported 78 AI export applications on July 6, 2026; designation and country delivery were not yet established by that announcement. DFC’s September IFC trade-finance counter-guaranty was board-approved, while its Orion consortium investment had closed in January. [2–3]')

d.add_page_break();head('What country records establish')
para('The new evidence ledger contains 2,170 source observations. The country-program matrix records 276 observed pairs across 18 inventory records and 121 country or territory codes. Its full comparison frame is 249 code units by 110 in-scope or candidate records. All other cells remain unknown. Coverage is concentrated in agencies that publish bulk financing and project records; it cannot be read as a distribution of all U.S. assistance or program intensity. [1]')
table(['Source and period','Verified extraction','Interpretation'],[
('EXIM\nFY2026 through June 30','908 non-cancelled authorization rows; 283 name 65 foreign country units. Another 550 cover multiple countries and 75 are U.S.-coded.','Country is the location of credit risk. Authorizations do not establish shipment, realized exports or program impact.'),
('DFC\nActive stock at June 30','1,085 portfolio rows, including 230 Legacy rows. The 855 other rows span 91 named country units, plus regional or undisclosed locations.','Active commitments can originate in earlier years. Rows include frameworks and repeated project numbers; summing them would risk double counting.'),
('MCC\nRoster retrieved October 15','97 lifecycle records: eight Implementation, five Signed, eight Development, 59 Closed, one Closure, 15 Terminated and one Discontinued.','Implementation is MCC’s reported stage. Signed and Development are pipeline; a terminated project is not a country-wide termination.'),
('Selected delivery reports\n2026','CLDP events in six named countries; USDA retrospective mission evidence in nine; Phoenix/FIRST evidence in four.','These are dated observations with source and geographic roles, not complete agency footprints or independent effectiveness evaluations.')
],[1.35,2.75,3.0])
para('The two DFC exports reconcile on country, department and commitment amount. The detailed file plus DFC’s Finance FAQ resolves 276 third-party lender guaranty rows and 38 certificates-of-participation rows into the guarantee and debt categories. The latter are direct funding despite “guaranty” in their name. Another 294 Finance rows remain unallocated without authoritative expansion of the DI code or disclosure of redacted/blank values. Legacy credit rows are kept separate from current named USAID initiatives. [4]')
para('Country location has different meanings across the evidence. EXIM’s credit-risk location, a DFC project site, an MCC compact country, an event venue and a foreign delegation’s origin are not equivalent. For example, NASDA reports a RAPP-supported Finnish delegation in Maine in June 2026. That supports reported delivery under an existing award, but neither overseas delivery in Finland nor Finland’s eligibility as a RAPP target market. [5]')
head('Eligibility and overseas presence remain separate')
para('The earlier 62-tool eligibility, sector and network matrices are retained without inferring new rules from activity. Their 15,438 country-tool cells include 12,436 unresolved eligibility entries. Agency service coverage does not establish a dedicated program office, embassy staffing or delivered assistance. New overlays receive no assumed country eligibility or office locations. A statistical similarity projection would presently mix unlike geographic concepts and extensive missing data. [1]')

d.add_page_break();head('Legacy initiatives and intake status')
table(['Initiative','Current finding'],[
('Phoenix and FIRST','Slovakia’s utility and the implementer confirm a completed SMR feasibility study in early 2026. The implementer places Phoenix under FIRST and names ongoing work in Estonia, Ukraine and Slovenia. Separate Phoenix intake and suboffice ownership remain unverified. [6]'),
('RAPP','USDA OIG’s June 2026 report records FAS’s decision not to award additional RAPP funding. Existing-award delivery is still reported. This resolves the distinction between new funding and continued implementation; it does not establish whole-program termination. [5]'),
('Prosper and Power Africa','ATI/PATI share a qualified contract crosswalk. A government court filing lists February 11–12, 2025 termination notices for Green Powered Technology, RTI EECA and Chemonics ESA. Future end dates do not verify continuation. Two legacy IDs now match federal records; two interagency action chains remain unverified. [7, 11]'),
('DCCP','CRDF explicitly reports DCCP funding for the Skopje accelerator. Commerce dates it June 10–11, 2026 and names State’s Bureau of Cyberspace and Digital Policy as convenor with CLDP. This verifies one component’s reported delivery; it does not resolve umbrella intake or surviving award IDs. [12]'),
('MDCP and STEP','MDCP announces no FY2027 competition; seven of 12 listed awards have elapsed end dates. STEP’s FY2026 federal competition closed August 18. SBA assigns national management to its Office of Manufacturing and Trade; state intake varies. Connecticut publishes a current route using FY2025 funds. [8]')
],[1.3,5.8])
para('ITSI’s ASU component names State EEB and publishes a license-request route. State’s 2026 EEB/CBA vacancy verifies commercial functions, but not ownership of BIDS, BFIF, Direct Line or the central Deal Team. Seven of 15 legacy records still lack a named program assignment; four have qualified roles and four have agency/contact verification. [1, 9–10]')
head('How this paper can be used')
para('This is a source-linked map of instruments and selected activity, not a complete active-program census or effectiveness ranking. The focused audit covers 25 selected legacy awards, including 13 Power Africa awards; none has certified current survival in this review. Historical notices, project-level cancellation reports and current continuation remain separate. The new DCCP country observation shares an event with CLDP and is not a second deal. Missing program ownership and comparable country evidence still prevent a defensible similarity visualization.')
head('Sources')
refs=[
'[1] Companion census, country ledger, methods and retained v0.7 award audit. All counts are reproducible; raw evidence and full URLs are indexed in the package.',
'[2] Commerce, American AI Exports application statement, July 6, 2026 (AI_APPLICATIONS); SABIT official program pages.',
'[3] DFC, investment announcements, February 4 and September 16, 2026 (DFC_MINERALS; DFC_SEP2026).',
'[4] EXIM FY2026 Q3 CSV and data dictionary; DFC June 30 active portfolio and Finance FAQ; MCC Where We Work lifecycle table. Agency URLs and snapshots accompany each extracted row.',
'[5] USDA OIG Audit 07601-0003-41, June 23, 2026, pp. 5–7; NASDA Inbound Trade Missions (RAPP_NASDA).',
'[6] Slovenské elektrárne, January 15, 2026; Sargent & Lundy, February 9, 2026 (PHOENIX_SEAS_CURRENT; PHOENIX_SL).',
'[7] DAI ATI page; retained PATI and DCCP Pacific award histories; indexed USAID ATI manual (PA00ZBRJ), contract 7200AA21C00056. Full manual retrieval failed; see qualified crosswalk.',
'[8] Commerce MDCP roster; SBA grants/OMT rosters; Grants.gov opportunity 363144, revised FY2026 STEP NOFO; Connecticut STEP FY2025 portal (V09_STEP_CT).',
'[9] GAO-26-108607, September 30, 2026, gao.gov/products/gao-26-108607.',
'[10] State vacancy E/EEB-MPI-2026-0021, August 12–18, 2026, usajobs.gov/job/880036200.',
'[11] DOJ, Global Health Council v. Trump, ECF 20 and 20-1, February 12, 2025, attachment pp. 2, 4, 11 (RECAP); USAspending FAIN/transaction records; DOGE API. Full IDs, dates, provenance and limits in v0.11 action review.',
'[12] CRDF Global, Skopje accelerator account; Commerce CLDP, June 10–11, 2026 event report (V10_CRDF_BALKANS; V10_CLDP_BALKANS). Full URLs and evidence notes in companion index.']
for t in refs:
 p=para(t);p.paragraph_format.space_after=Pt(2)
 for r in p.runs:r.font.size=Pt(8)
for style in d.styles:
 for border in list(style.element.xpath('.//w:pBdr')):border.getparent().remove(border)
for p in d.paragraphs:
 for border in list(p._p.xpath('.//w:pBdr')):border.getparent().remove(border)
d.save(O/'US_Trade_Facilitation_Informative_Paper_v0_11.docx')
print('Saved paper')

